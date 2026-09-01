"""Build legal-style .docx bytes from clause cards, playbooks, redlines, and drafts."""

from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

from app.contracts.redline import RedlineOp
from app.schemas.contracts import (
    ClauseExtractionResponse,
    ClauseStatus,
    PlaybookReviewResponse,
    PlaybookVerdict,
    RedlineHunkResponse,
    RedlineResponse,
    ReviewTableResponse,
)

_NAVY = RGBColor(0x1B, 0x3A, 0x4B)
_RED = RGBColor(0x9B, 0x1C, 0x1C)
_BLUE = RGBColor(0x1D, 0x4E, 0x89)


def export_clauses(extraction: ClauseExtractionResponse) -> bytes:
    doc = _base_document("Clause extraction")
    doc.add_heading(extraction.filename, level=2)
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    _header_row(table, ["Term", "Status", "Value", "Quote"])
    for card in extraction.cards:
        cells = table.add_row().cells
        cells[0].text = card.label
        cells[1].text = card.status.value
        cells[2].text = card.value or "—"
        cells[3].text = card.quote or "—"
    _footer(doc)
    return _save(doc)


def export_playbook(review: PlaybookReviewResponse) -> bytes:
    doc = _base_document("Playbook review")
    subtitle = doc.add_paragraph()
    run = subtitle.add_run(
        f"{review.playbook_name}  ·  {review.filename}"
    )
    run.italic = True
    summary = review.summary
    doc.add_paragraph(
        f"Pass {summary.pass_count}  ·  Fallback {summary.fallback_count}  ·  "
        f"Missing {summary.missing_count}  ·  Off-playbook {summary.off_playbook_count}  ·  "
        f"Blockers {summary.blocker_count}"
    )
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    _header_row(table, ["Term", "Verdict", "Extracted", "Suggested language"])
    for finding in review.findings:
        cells = table.add_row().cells
        cells[0].text = finding.label
        cells[1].text = finding.verdict.value.replace("_", " ")
        cells[2].text = finding.extracted_value or "—"
        cells[3].text = finding.suggested_language or "—"
        if finding.verdict in {
            PlaybookVerdict.MISSING,
            PlaybookVerdict.OFF_PLAYBOOK,
        }:
            for cell in cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = _RED
    doc.add_heading("Notes", level=2)
    for finding in review.findings:
        if finding.verdict in {
            PlaybookVerdict.PASS,
            PlaybookVerdict.NOT_APPLICABLE,
        }:
            continue
        doc.add_paragraph(
            f"{finding.label}: {finding.reason}",
            style="List Bullet",
        )
    _footer(doc)
    return _save(doc)


def export_redline(redline: RedlineResponse) -> bytes:
    title = (
        "Playbook redline"
        if redline.mode == "playbook"
        else "Document redline"
    )
    doc = _base_document(title)
    caption = doc.add_paragraph()
    caption.add_run(
        f"{redline.left_filename or 'Current'}  →  "
        f"{redline.right_filename or 'Proposed'}"
    ).italic = True
    if redline.unchanged:
        doc.add_paragraph("No differences.")
        _footer(doc)
        return _save(doc)
    for hunk in redline.hunks:
        doc.add_heading(hunk.label, level=2)
        paragraph = doc.add_paragraph()
        _write_spans(paragraph, hunk)
    _footer(doc)
    return _save(doc)


def export_review_table(table: ReviewTableResponse) -> bytes:
    doc = _base_document("Review table")
    headers = ["Document"] + [column.label for column in table.columns]
    grid = doc.add_table(rows=1, cols=len(headers))
    grid.style = "Table Grid"
    _header_row(grid, headers)
    for row in table.rows:
        cells = {cell.key: cell for cell in row.cells}
        values = [row.filename]
        for column in table.columns:
            cell = cells.get(column.key)
            if cell is None or cell.status == ClauseStatus.MISSING:
                values.append("—")
            else:
                values.append(cell.value or cell.status.value)
        added = grid.add_row().cells
        for index, value in enumerate(values):
            added[index].text = value[:400]
    _footer(doc)
    return _save(doc)


def export_draft(*, title: str, body: str) -> bytes:
    doc = _base_document(title.strip() or "Legal draft")
    for block in (body or "").split("\n"):
        line = block.rstrip()
        if not line:
            doc.add_paragraph("")
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            paragraph = doc.add_paragraph(line)
            paragraph.paragraph_format.space_after = Pt(8)
    _footer(doc)
    return _save(doc)


def _base_document(title: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    heading = doc.add_heading(title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in heading.runs:
        run.font.color.rgb = _NAVY
        run.font.name = "Times New Roman"
    return doc


def _header_row(table, labels: list[str]) -> None:
    for index, label in enumerate(labels):
        cell = table.rows[0].cells[index]
        cell.text = label
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
        _shade(cell, "E8EEF2")


def _shade(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")
    tc_pr.append(shading)


def _write_spans(paragraph, hunk: RedlineHunkResponse) -> None:
    if not hunk.spans:
        paragraph.add_run(hunk.right or hunk.left or "")
        return
    for span in hunk.spans:
        run = paragraph.add_run(span.text)
        run.font.name = "Times New Roman"
        if span.op == RedlineOp.DELETE.value:
            run.font.strike = True
            run.font.color.rgb = _RED
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif span.op == RedlineOp.INSERT.value:
            run.underline = True
            run.font.color.rgb = _BLUE
            run.bold = True


def _footer(doc: Document) -> None:
    note = doc.add_paragraph()
    run = note.add_run(
        "Generated by LegalGPT. This is a working draft, not legal advice. "
        "Verify against the source document before sending."
    )
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def _save(doc: Document) -> bytes:
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
